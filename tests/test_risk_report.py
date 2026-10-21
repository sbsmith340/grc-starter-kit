import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import risk_report as rr  # noqa: E402

ROWS = rr.load(ROOT / "register" / "risk_register.csv")
TODAY = date(2026, 10, 21)


def base(**kw):
    r = dict(ROWS[0])
    r.update(kw)
    return r


def test_sample_register_is_valid():
    assert rr.validate(ROWS, TODAY) == []


def test_ratings_bands():
    assert [rr.rating(s) for s in (25, 20, 19, 12, 11, 6, 5, 1)] == \
        ["Critical", "Critical", "High", "High", "Medium", "Medium", "Low", "Low"]


def test_residual_above_inherent_rejected():
    errs = rr.validate([base(Likelihood="2", Impact="2", ResidualLikelihood="3", ResidualImpact="3")], TODAY)
    assert any("higher than inherent" in e for e in errs)


def test_out_of_range_and_bad_cia_rejected():
    errs = rr.validate([base(Likelihood="6", Confidentiality="Severe")], TODAY)
    assert any("Likelihood must be 1-5" in e for e in errs)
    assert any("Confidentiality" in e for e in errs)


def test_accepted_risk_needs_unexpired_expiry():
    assert any("needs an AcceptanceExpiry" in e for e in rr.validate([base(Treatment="Accept", AcceptanceExpiry="")], TODAY))
    assert any("expired" in e for e in rr.validate([base(Treatment="Accept", AcceptanceExpiry="2026-01-01")], TODAY))


def test_duplicate_ids_rejected():
    assert any("duplicate" in e for e in rr.validate([ROWS[0], ROWS[0]], TODAY))


def test_report_flags_risks_above_appetite():
    over = base(RiskId="R-099", Likelihood="5", Impact="5", ResidualLikelihood="4", ResidualImpact="4")
    section = rr.report(ROWS + [over]).split("## Above appetite after treatment")[1]
    assert "R-099" in section and "residual 16 (High)" in section
    assert "None." in rr.report(ROWS).split("## Above appetite after treatment")[1]


def test_heatmap_places_ids():
    hm = rr.heatmap(ROWS, "Likelihood", "Impact")
    row4 = next(l for l in hm.splitlines() if l.startswith("| **4**"))
    assert "R-001" in row4 and "R-002" in row4


def test_workbook_builds_with_all_tabs(tmp_path, monkeypatch):
    sys.path.insert(0, str(ROOT / "inventory"))
    import build_workbook as bw
    from openpyxl import load_workbook
    monkeypatch.setattr(bw, "OUT", tmp_path / "inv.xlsx")
    wb = load_workbook(bw.build())
    assert wb.sheetnames == ["How to collect"] + list(bw.TABS)
    assert wb["Hardware"]["A1"].value == "Device Name"
    assert wb["Hardware"].freeze_panes == "A2"
