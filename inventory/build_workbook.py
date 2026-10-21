"""
Build the asset inventory workbook template.

  python inventory/build_workbook.py      # -> inventory/asset-inventory-template.xlsx

One tab per asset class, with header notes, dropdowns for controlled values,
frozen headers, and a "How to collect" tab pointing at the query for each tab
(see collection-queries.md).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("asset-inventory-template.xlsx")
YES_NO = ["Yes", "No", "Unknown"]
CIA = ["Low", "Moderate", "High"]

# (column, width, note, allowed values or None)
TABS = {
    "Hardware": [
        ("Device Name", 22, "Hostname as it appears in Intune / AD / EDR", None),
        ("Device Type", 14, None, ["Laptop", "Desktop", "Server", "VM", "Mobile", "Network", "Printer/MFD", "IoT"]),
        ("Manufacturer / Model", 22, None, None),
        ("Serial Number", 18, None, None),
        ("Owner / Assigned User", 24, None, None),
        ("Environment", 12, None, ["Production", "Staging", "Dev", "Lab"]),
        ("OS & Version", 22, None, None),
        ("Support Status", 14, "From Microsoft / vendor lifecycle", ["Supported", "Expiring", "Unsupported"]),
        ("Disk Encryption", 12, None, YES_NO),
        ("EDR Onboarded", 12, None, YES_NO),
        ("Internet Exposed", 12, "Any inbound path from the internet", YES_NO),
        ("Highest Data Classification", 16, "Per the data classification standard", CIA),
        ("Last Check-in", 14, "Stale > 30 days = investigate", None),
        ("Source", 16, "Intune / Defender / AD / manual", None),
    ],
    "Software": [
        ("Software Name", 26, None, None),
        ("Vendor", 18, None, None),
        ("Version(s)", 14, None, None),
        ("Install Count", 12, None, None),
        ("Approved", 10, "On the approved software list?", YES_NO),
        ("Business Owner", 20, None, None),
        ("Data Types Accessed", 26, "Of the eight classification types", None),
        ("Vendor Supported", 14, None, YES_NO),
        ("Known Critical CVEs", 14, "From Defender Vulnerability Management", None),
        ("Source", 16, None, None),
    ],
    "Cloud": [
        ("Provider", 12, None, ["Azure", "AWS", "GCP", "SaaS", "Other"]),
        ("Account / Subscription", 26, None, None),
        ("Resource / Service", 26, None, None),
        ("Region", 12, None, None),
        ("Publicly Accessible", 12, "Public IP, public bucket/container, or anonymous access", YES_NO),
        ("Encryption at Rest", 12, None, YES_NO),
        ("Logging to SIEM", 12, None, YES_NO),
        ("Highest Data Classification", 16, None, CIA),
        ("Owner", 20, None, None),
    ],
    "External Domains": [
        ("Domain / Hostname", 28, None, None),
        ("Resolves To (IP / CNAME)", 26, None, None),
        ("Hosting Provider", 18, None, None),
        ("Purpose", 22, None, None),
        ("TLS Valid / Expiry", 16, None, None),
        ("SPF / DKIM / DMARC", 18, "For mail-sending domains", None),
        ("Registrar & Expiry", 18, "Expired domains get hijacked", None),
        ("Owner", 18, None, None),
    ],
    "External Authentication": [
        ("Login Endpoint / App", 30, "Every place a user can sign in from the internet", None),
        ("Identity Provider", 16, None, ["Entra ID SSO", "Local accounts", "Other IdP"]),
        ("MFA Enforced", 12, None, YES_NO),
        ("Phishing-resistant MFA", 14, "FIDO2 / passkey / cert-based", YES_NO),
        ("Legacy Auth Blocked", 12, None, YES_NO),
        ("Lockout Threshold", 14, None, None),
        ("Users With Access", 12, None, None),
        ("Owner", 18, None, None),
    ],
    "Third Parties": [
        ("Vendor", 22, None, None),
        ("Service Provided", 26, None, None),
        ("Data Shared (types)", 26, "Of the eight classification types", None),
        ("Access to PII / PHI", 12, None, YES_NO),
        ("Integration / Access Method", 22, "SSO app, API key, VPN, shared account...", None),
        ("Assurance Evidence", 20, "SOC 2 Type II, ISO 27001 cert, questionnaire", None),
        ("Contract Has Breach Notice", 14, None, YES_NO),
        ("Criticality", 12, None, ["Low", "Moderate", "High"]),
        ("Next Review", 12, None, None),
        ("Owner", 18, None, None),
    ],
}

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def build() -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols in TABS.items():
        ws = wb.create_sheet(name)
        for idx, (col, width, note, allowed) in enumerate(cols, start=1):
            c = ws.cell(row=1, column=idx, value=col)
            c.fill, c.font = HEADER_FILL, HEADER_FONT
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(idx)].width = width
            if note:
                c.comment = Comment(note, "GRC starter kit")
            if allowed:
                dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=True)
                dv.error, dv.errorTitle = f"Choose one of: {', '.join(allowed)}", "Invalid value"
                ws.add_data_validation(dv)
                dv.add(f"{get_column_letter(idx)}2:{get_column_letter(idx)}2000")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
        ws.row_dimensions[1].height = 32

    how = wb.create_sheet("How to collect", 0)
    how["A1"] = "Asset inventory: how to collect each tab"
    how["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Tab", "Primary source", "Query / command (see collection-queries.md)"),
        ("Hardware", "Intune + Defender for Endpoint + AD", "Defender Advanced Hunting: DeviceInfo; Graph: managedDevices"),
        ("Software", "Defender Vulnerability Management", "Advanced Hunting: DeviceTvmSoftwareInventory"),
        ("Cloud", "Azure Resource Graph, AWS Config", "az graph query; aws resourcegroupstaggingapi"),
        ("External Domains", "DNS zones, registrar, CT logs", "Azure DNS / Route 53 export; crt.sh lookup"),
        ("External Authentication", "Entra ID enterprise apps, sign-in logs", "Graph: servicePrincipals; KQL on SigninLogs"),
        ("Third Parties", "Procurement / AP, enterprise apps", "Accounts payable vendor list + SSO app list"),
        ("", "", ""),
        ("Rules", "", ""),
        ("1", "One row per asset; Source column says where the row came from", ""),
        ("2", "Unknown is an allowed answer - it is a finding, not a blank", ""),
        ("3", "Remove duplicates after merging sources (match on serial / hostname)", ""),
        ("4", "Review quarterly; stale Last Check-in > 30 days = investigate or retire", ""),
    ]
    for r, row in enumerate(rows, start=3):
        for c, v in enumerate(row, start=1):
            cell = how.cell(row=r, column=c, value=v)
            if r == 3 or v == "Rules":
                cell.font = Font(bold=True)
    for col, w in zip("ABC", (24, 44, 64)):
        how.column_dimensions[col].width = w
    wb.save(OUT)
    return OUT


if __name__ == "__main__":
    print(f"wrote {build()}")
